"""Read-only source inspection; write an auditable snapshot under results/q1."""
from pathlib import Path
import json, hashlib
import openpyxl
import rasterio
from pypdf import PdfReader
from docx import Document

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'results' / 'q1'

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    snapshot = {'files': {}, 'workbooks': {}, 'pdfs': {}, 'dem': {}}
    for p in (ROOT / '数据').rglob('*'):
        if p.is_file():
            snapshot['files'][str(p.relative_to(ROOT))] = {
                'bytes': p.stat().st_size, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
        if p.suffix == '.xlsx':
            wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
            sheets = {}
            for ws in wb:
                rows = list(ws.values)
                sheets[ws.title] = {'rows': rows, 'max_row': ws.max_row, 'max_column': ws.max_column}
            snapshot['workbooks'][p.name] = sheets
            wb.close()
        elif p.suffix == '.pdf':
            snapshot['pdfs'][p.name] = '\n'.join(page.extract_text() for page in PdfReader(p).pages)
        elif p.suffix == '.tif':
            with rasterio.open(p) as ds:
                a = ds.read(1, masked=True)
                snapshot['dem'] = {'file': str(p.relative_to(ROOT)), 'crs': str(ds.crs),
                    'shape': a.shape, 'transform': list(ds.transform), 'bounds': list(ds.bounds),
                    'nodata': ds.nodata, 'min': float(a.min()), 'max': float(a.max()),
                    'masked_count': int(a.mask.sum()) if hasattr(a.mask, 'sum') else 0}
    p = next(ROOT.glob('D-*.docx'))
    doc = Document(p)
    snapshot['problem_text'] = '\n'.join(x.text for x in doc.paragraphs)
    snapshot['files'][p.name] = {'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
    for p in (ROOT / '数据').rglob('*.csv'):
        data = p.read_bytes()
        try: text = data.decode('utf-8-sig')
        except UnicodeDecodeError: text = data.decode('gb18030')
        lines = text.splitlines()
        snapshot['files'][str(p.relative_to(ROOT))].update(lines=len(lines), first_lines=lines[:3])
    wb = openpyxl.load_workbook(ROOT / '结果提交模板.xlsx', read_only=True)
    snapshot['template'] = {ws.title: {'rows':ws.max_row,'cols':ws.max_column,'header':list(next(ws.values))} for ws in wb}
    wb.close()
    (OUT / 'input_inventory.json').write_text(json.dumps(snapshot,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'dem':snapshot['dem'],'pdfs':snapshot['pdfs'],'template':snapshot['template'],
          'workbooks':{k:{s:(d['max_row'],d['max_column']) for s,d in v.items()} for k,v in snapshot['workbooks'].items()}},ensure_ascii=False,indent=2))

if __name__ == '__main__':
    main()
