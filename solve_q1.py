"""Exact type-count batching for problem 1. Run from the project root.

Energy follows the explicitly documented interpretation in the model contract.
No aircraft, battery scheduling, deadline or communication constraints in Q1.
"""
from pathlib import Path
from functools import lru_cache
from itertools import product
from collections import Counter, defaultdict
import argparse, csv, json, math, time, hashlib, sys
import numpy as np
import openpyxl
import rasterio
from pyproj import Geod
from scipy.optimize import milp, LinearConstraint, Bounds

ROOT=Path(__file__).resolve().parent
DATA=ROOT/'数据'/'无人机应急物资运输基础数据'
OUT=ROOT/'results'/'q1'
KINDS=('MED','WAT','FOD','HYG')
GRAVITY=9.81
TOL=1e-9

def dump(name, data):
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')

def csvout(name, rows):
    if not rows:return
    with (OUT/name).open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)

def read_inputs():
    def rows(name,sheet='数据'):
        wb=openpyxl.load_workbook(DATA/name,data_only=True,read_only=True)
        result=list(wb[sheet].values);wb.close();return result
    nr=rows('调度中心与服务区.xlsx')
    nodes={r[0]:dict(id=r[0],name=r[1],lon=float(r[2]),lat=float(r[3]),z=float(r[4])) for r in [nr[2],*nr[6:21]]}
    fields=('id','name','empty_mass','capacity','volume','speed','range_empty','range_full',
            'battery','reserve_pct','prepare','load_per_box','handover','handover_per_box','up_speed','down_speed','eta','descent_eta')
    models={r[0]:dict(zip(fields,r)) for r in rows('运输无人机数据.xlsx')[2:5]}
    br=rows('物资需求与配送时限.xlsx','逐箱货箱清单')
    assert len(br)==81 and br[0][0]=='货箱编号'
    boxes=[dict(id=r[0],area=r[1],kind=r[0].split('-')[1],mass=float(r[3]),volume=float(r[4]),
                first=r[5]=='是',deadline=r[6],desired=r[7]) for r in br[1:]]
    assert len(set(b['id'] for b in boxes))==80
    assert boxes[0]['id']=='S001-MED-01' and boxes[-1]['id']=='S015-FOD-01'
    assert len(nodes)==16 and len(models)==3
    assert all(g['reserve_pct']==20 for g in models.values())
    for b in boxes:assert b['area'] in nodes and b['mass']>0 and b['volume']>0
    names={'医疗物资':'MED','饮用水':'WAT','应急食品':'FOD','生活卫生用品':'HYG'}
    counts=Counter((b['area'],b['kind']) for b in boxes)
    for r in rows('物资需求与配送时限.xlsx')[1:]:
        assert counts[(r[0],names[r[1]])]==r[2]
    attrs={k:next((b['mass'],b['volume']) for b in boxes if b['kind']==k) for k in KINDS}
    for b in boxes:assert (b['mass'],b['volume'])==attrs[b['kind']]
    return nodes,models,boxes,attrs

def crossed_cells(start,end,transform,shape):
    """All pixels intersecting the closed segment, including boundary touches."""
    c0,r0=(~transform)*start;c1,r1=(~transform)*end
    dc,dr=c1-c0,r1-r0
    ts=[0.,1.]
    for p,q in ((c0,c1),(r0,r1)):
        if abs(q-p)>1e-14:
            for k in range(math.ceil(min(p,q)),math.floor(max(p,q))+1):
                t=(k-p)/(q-p)
                if 0<t<1:ts.append(t)
    ts=sorted(set(ts));cells=set()
    def indices(x):
        k=round(x)
        return (k-1,k) if abs(x-k)<1e-9 else (math.floor(x),)
    for t in [*ts,*[(a+b)/2 for a,b in zip(ts[:-1],ts[1:])]]:
        for r,c in product(indices(r0+dr*t),indices(c0+dc*t)):
            if not (0<=r<shape[0] and 0<=c<shape[1]):raise ValueError('Route outside DEM')
            cells.add((r,c))
    return sorted(cells)

def geometry(nodes):
    dempath=next((ROOT/'数据').rglob('*.tif'))
    geod=Geod(ellps='WGS84');home=nodes['O01'];routes={}
    with rasterio.open(dempath) as ds:
        grid=ds.read(1);tr=ds.transform
        assert str(ds.crs)=='EPSG:4326' and np.isfinite(grid).all()
        for sid,s in nodes.items():
            if sid=='O01':continue
            start=(home['lon'],home['lat']);end=(s['lon'],s['lat'])
            cells=crossed_cells(start,end,tr,grid.shape)
            assert cells==crossed_cells(end,start,tr,grid.shape)
            terrain=np.array([grid[r,c] for r,c in cells],dtype=float)
            assert np.all(terrain!=-32767)
            high=float(terrain.max());H=high+50
            up=H-home['z'];down=H-s['z']-30
            assert up>=0 and down>=0
            _,_,distance=geod.inv(*start,*end)
            t=np.linspace(0,1,math.ceil(distance)+1)
            cc,rr=(~tr)*(start[0]+t*(end[0]-start[0]),start[1]+t*(end[1]-start[1]))
            profile=grid[np.floor(rr).astype(int),np.floor(cc).astype(int)].astype(float)
            assert profile.max()<=high+TOL
            # Local equirectangular WGS84 ellipsoidal approximation.
            phi=math.radians((start[1]+end[1])/2);e2=6.69437999014e-3;a=6378137.
            M=a*(1-e2)/(1-e2*math.sin(phi)**2)**1.5;Nu=a/math.sqrt(1-e2*math.sin(phi)**2)
            local=math.hypot(M*math.radians(end[1]-start[1]),Nu*math.cos(phi)*math.radians(end[0]-start[0]))
            routes[sid]=dict(area=sid,distance_m=distance,max_terrain_m=high,cruise_m=H,
                outbound_up_m=up,return_up_m=down,cells=len(cells),fine_sample_max_m=float(profile.max()),
                local_distance_m=local,relative_distance_difference=abs(local-distance)/distance,
                profile_distance_m=(t*distance)[::10].tolist(),profile_terrain_m=profile[::10].tolist())
    return routes

def energy(route,g,q):
    assert -TOL<=q<=g['capacity']+TOL
    L=g['range_empty']-(g['range_empty']-g['range_full'])*(q/g['capacity'])**1.5
    d=route['distance_m'];up=route['outbound_up_m'];ret=route['return_up_m']
    hor=g['battery']*d*(1/L+1/g['range_empty'])
    climb=GRAVITY*((g['empty_mass']+q)*up+g['empty_mass']*ret)/(3.6e6*g['eta'])
    return hor+climb

def safe_load(route,g,rho):
    budget=(1-rho)*g['battery'];Q=g['capacity']
    if energy(route,g,0)>budget:return None
    if energy(route,g,Q)<=budget:return float(Q)
    lo,hi=0.,float(Q)
    while hi-lo>1e-8:
        mid=(lo+hi)/2
        if energy(route,g,mid)<=budget:lo=mid
        else:hi=mid
    assert energy(route,g,lo)<=budget+TOL
    assert energy(route,g,min(Q,lo+1e-6))>budget
    return lo

def times(route,g,n):
    fly=2*route['distance_m']/g['speed']+(route['outbound_up_m']+route['return_up_m'])*(1/g['up_speed']+1/g['down_speed'])
    back=fly+g['handover']+n*g['handover_per_box']
    return fly,back,back+g['prepare']+n*g['load_per_box']

def patterns(route,models,boxes,attrs,rho):
    demands=tuple(sum(b['kind']==k for b in boxes) for k in KINDS);out=[]
    for count in product(*(range(n+1) for n in demands)):
        n=sum(count)
        if n==0:continue
        q=sum(c*attrs[k][0] for c,k in zip(count,KINDS));vol=sum(c*attrs[k][1] for c,k in zip(count,KINDS))
        for gid,g in models.items():
            if q>g['capacity']+TOL or vol>g['volume']+TOL:continue
            E=energy(route,g,q)
            if E>(1-rho)*g['battery']+TOL:continue
            fly,back,work=times(route,g,n)
            out.append(dict(model=gid,count=count,n=n,mass_kg=q,volume_m3=vol,energy_kwh=E,
                            flight_s=fly,roundtrip_s=back,work_s=work,soc_pct=100*(1-E/g['battery'])))
    return demands,out

def optimize(demands,pats,order=(0,1,2)):
    preds={};visited=[]
    @lru_cache(None)
    def dp(state):
        if not any(state):return (0.,0.,0.)
        best=None;pick=None
        for j,p in enumerate(pats):
            if any(a>b for a,b in zip(p['count'],state)):continue
            rest=tuple(b-a for a,b in zip(p['count'],state));old=dp(rest)
            if old is None:continue
            val=(old[0]+1,old[1]+p['energy_kwh'],old[2]+p['work_s'])
            if best is None or tuple(val[k] for k in order)<tuple(best[k] for k in order):best=val;pick=(rest,j)
        preds[state]=pick
        visited.append(dict(boxes=sum(state),state=list(state),objectives=best))
        return best
    optimum=dp(demands)
    if optimum is None:return None,[],visited
    state=demands;seq=[]
    while any(state):
        state,j=preds[state];seq.append(pats[j].copy())
    return optimum,seq,visited

def allocate(sid,seq,boxes,start=1):
    queues={k:sorted(b['id'] for b in boxes if b['kind']==k) for k in KINDS}
    records=[]
    for j,p in enumerate(seq,start):
        ids=[]
        for k,n in zip(KINDS,p['count']):ids.extend(queues[k][:n]);queues[k]=queues[k][n:]
        records.append(dict(sortie=f'Q1-{j:03d}',area=sid,model=p['model'],boxes=';'.join(ids),
            n=p['n'],mass_kg=p['mass_kg'],volume_m3=p['volume_m3'],flight_s=p['flight_s'],
            roundtrip_s=p['roundtrip_s'],work_s=p['work_s'],energy_kwh=p['energy_kwh'],soc_pct=p['soc_pct']))
    assert not any(queues.values())
    return records

def validate(records,boxes,models,routes,rho):
    lookup={b['id']:b for b in boxes};seen=[]
    for p in records:
        ids=p['boxes'].split(';');seen.extend(ids);bs=[lookup[b] for b in ids];g=models[p['model']]
        assert all(b['area']==p['area'] for b in bs)
        q=sum(b['mass'] for b in bs);v=sum(b['volume'] for b in bs)
        assert abs(q-p['mass_kg'])<TOL and abs(v-p['volume_m3'])<TOL
        assert q<=g['capacity']+TOL and v<=g['volume']+TOL
        E=energy(routes[p['area']],g,q)
        assert abs(E-p['energy_kwh'])<TOL and E<=(1-rho)*g['battery']+TOL
        assert abs(p['soc_pct']-100*(1-E/g['battery']))<1e-7
        for key,value in zip(('flight_s','roundtrip_s','work_s'),times(routes[p['area']],g,len(ids))):
            assert abs(p[key]-value)<1e-7
        assert p['soc_pct']>=rho*100-1e-7
    assert Counter(seen)==Counter(lookup.keys())
    return dict(boxes=len(seen),unique=len(set(seen)),mass_kg=sum(p['mass_kg'] for p in records),
                min_soc_pct=min(p['soc_pct'] for p in records),constraint_violations=0)

def milp_check(demand,pats,opt):
    A=np.array([p['count'] for p in pats],dtype=float).T
    ints=np.ones(len(pats));bounds=Bounds(0,np.inf)
    base=LinearConstraint(A,np.array(demand),np.array(demand))
    first=milp(np.ones(len(pats)),integrality=ints,bounds=bounds,constraints=base,options={'mip_rel_gap':0.0,'time_limit':60})
    assert first.success and abs(first.fun-opt[0])<1e-6
    second=milp(np.array([p['energy_kwh'] for p in pats]),integrality=ints,bounds=bounds,
        constraints=[base,LinearConstraint(np.ones((1,len(pats))),opt[0],opt[0])],options={'mip_rel_gap':0.0,'time_limit':60})
    assert second.success and abs(second.fun-opt[1])<1e-6
    return dict(N=int(round(first.fun)),E=second.fun,energy_difference=abs(second.fun-opt[1]),status_N=first.status,status_E=second.status)

def run(minimal=False):
    start=time.perf_counter();nodes,models,boxes,attrs=read_inputs();routes=geometry(nodes)
    rho=.2;areas=['S001'] if minimal else sorted(routes)
    schedules=[];caps=[];pattern_stats=[];certs={};dp_traces={}
    policies={'NET':(0,1,2),'ENT':(1,0,2),'TNE':(2,0,1),'NTE':(0,2,1)}
    comparisons={key:[] for key in policies}
    for sid in areas:
        bs=[b for b in boxes if b['area']==sid];route=routes[sid]
        for gid,g in models.items():
            cap=safe_load(route,g,rho)
            caps.append(dict(area=sid,model=gid,max_safe_kg=cap,rated_kg=g['capacity'],rho=rho,
                             empty_soc_pct=100*(1-energy(route,g,0)/g['battery']),
                             full_soc_pct=100*(1-energy(route,g,g['capacity'])/g['battery'])))
        demand,pats=patterns(route,models,bs,attrs,rho)
        opt,seq,trace=optimize(demand,pats);assert opt is not None
        certs[sid]=milp_check(demand,pats,opt)
        dp_traces[sid]=trace
        schedules.extend(allocate(sid,seq,bs,len(schedules)+1))
        pattern_stats.append(dict(area=sid,states=math.prod(n+1 for n in demand),
            patterns=len(pats),A=sum(p['model']=='A' for p in pats),B=sum(p['model']=='B' for p in pats),C=sum(p['model']=='C' for p in pats)))
        for key,order in policies.items():
            val,ss,_=optimize(demand,pats,order)
            comparisons[key].extend(allocate(sid,ss,bs,len(comparisons[key])+1))
    selected=[b for b in boxes if b['area'] in areas]
    checks=validate(schedules,selected,models,routes,rho)
    for records in comparisons.values():validate(records,selected,models,routes,rho)
    totals=dict(N=len(schedules),E=sum(p['energy_kwh'] for p in schedules),T=sum(p['work_s'] for p in schedules),
                flight_s=sum(p['flight_s'] for p in schedules),roundtrip_s=sum(p['roundtrip_s'] for p in schedules))
    result=dict(totals=totals,checks=checks,capacities=caps,schedule=schedules,geometry=routes,models=models,
                nodes=nodes,pattern_stats=pattern_stats,milp_certificates=certs,policies={
        k:dict(N=len(v),E=sum(p['energy_kwh'] for p in v),T=sum(p['work_s'] for p in v),schedule=v)
        for k,v in comparisons.items()},runtime_s=time.perf_counter()-start)
    if minimal:
        dump('minimal.json',result);print(json.dumps(dict(mode='minimal',**totals,checks=checks),ensure_ascii=False));return
    sens=[];scaps=[]
    for percent in range(51):
        reserve=percent/100;ss=[];feasible=True
        for sid in areas:
            bs=[b for b in boxes if b['area']==sid]
            for gid,g in models.items():scaps.append(dict(rho=reserve,area=sid,model=gid,max_safe_kg=safe_load(routes[sid],g,reserve)))
            dd,pp=patterns(routes[sid],models,bs,attrs,reserve);vv,seq,_=optimize(dd,pp)
            if vv is None:feasible=False;continue
            ss.extend(allocate(sid,seq,bs,len(ss)+1))
        if feasible:validate(ss,boxes,models,routes,reserve)
        sens.append(dict(rho=reserve,feasible=feasible,N=len(ss) if feasible else None,
                         E=sum(p['energy_kwh'] for p in ss) if feasible else None,T=sum(p['work_s'] for p in ss) if feasible else None))
    for sid in areas:
        for gid in models:
            seq=[c['max_safe_kg'] for c in scaps if c['area']==sid and c['model']==gid]
            assert all(b is None or (a is not None and b<=a+TOL) for a,b in zip(seq[:-1],seq[1:]))
    ns=[s['N'] for s in sens if s['feasible']];assert all(b>=a for a,b in zip(ns[:-1],ns[1:]))
    result['sensitivity']=sens;result['sensitivity_capacities']=scaps;result['runtime_s']=time.perf_counter()-start
    dump('solution.json',result);dump('dp_trace.json',dp_traces)
    csvout('max_safe_payload.csv',caps);csvout('sorties.csv',schedules);csvout('sensitivity.csv',sens);csvout('sensitivity_payload.csv',scaps)
    csvout('geometry.csv',[{k:v for k,v in r.items() if not k.startswith('profile')} for r in routes.values()])
    csvout('pattern_counts.csv',pattern_stats);csvout('boxes.csv',boxes)
    csvout('policy_comparison.csv',[dict(policy=k,**{a:b for a,b in v.items() if a!='schedule'}) for k,v in result['policies'].items()])
    wb=openpyxl.load_workbook(ROOT/'结果提交模板.xlsx');ws=wb['Q1_单点组批']
    for i,p in enumerate(schedules,2):
        vals=[p[k] for k in ('sortie','area','model','boxes','mass_kg','volume_m3','roundtrip_s','energy_kwh','soc_pct')]
        for j,v in enumerate(vals,1):ws.cell(i,j,v)
    from openpyxl.comments import Comment
    ws['G1'].comment=Comment('往返时间=去返飞行+服务区交接，不含起飞前准备和装载；累计作业时间另见sorties.csv的work_s。','Q1')
    wb.save(ROOT/'results'/'问题1_结果.xlsx')
    print(json.dumps(dict(mode='full',**totals,checks=checks,policies={k:{a:b for a,b in v.items() if a!='schedule'} for k,v in result['policies'].items()},runtime_s=result['runtime_s']),ensure_ascii=False,indent=2))

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--minimal',action='store_true');args=ap.parse_args()
    run(args.minimal)
